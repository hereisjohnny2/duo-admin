"""
Importa os lançamentos das abas mensais (JAN..OUT) da planilha CTR03.R00_FINANCEIRO_2026.xlsx
para uma migration SQL que popula as tabelas `categorias` e `contas`.

Uso:
  python scripts/import-contas-xlsx.py <caminho-da-planilha> <caminho-sql-saida>
"""
import sys
import datetime
import openpyxl

MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT"]

STATUS_MAP = {
    "PAGO": "Pago",
    "EM ABERTO": "Em Aberto",
    "ATRASADO": "Atrasado",
}

# Mapeia o texto (em caixa alta) da planilha para o nome já seedado em
# 20260923120000_add_contas.sql (mesma grafia/acentuação, casing legível).
CATEGORIA_MAP = {
    "ÁGUA": "Água",
    "ANUIDADE CARTÃO": "Anuidade Cartão",
    "CAIXINHA": "Caixinha",
    "CELULAR": "Celular",
    "CONDOMÍNIO": "Condomínio",
    "ENERGIA": "Energia",
    "INTERNET": "Internet",
    "IPTU": "IPTU",
    "IPVA": "IPVA",
    "LIMPEZA": "Limpeza",
    "MULTA DARF": "Multa DARF",
    "OUTROS": "Outros",
    "PLANO DE SAÚDE": "Plano de Saúde",
    "PROLABORE": "Prolabore",
    "SALÁRIO": "Salário",
    "SIMPLES": "Simples",
}


def normalize_categoria(raw: str) -> str:
    return CATEGORIA_MAP.get(raw.strip().upper(), raw.strip())


def esc(s: str) -> str:
    return s.replace("'", "''")


def to_iso_date(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = str(v).strip()
    if s in ("", "-", "?"):
        return ""
    # tenta "dd/mm/yyyy" ou "dd//mm/yyyy" (erro de digitação na planilha)
    s2 = s.replace("//", "/")
    parts = s2.split("/")
    if len(parts) == 3:
        try:
            d, m, y = (int(p) for p in parts)
            return datetime.date(y, m, d).isoformat()
        except ValueError:
            pass
    return ""  # texto livre não reconhecido como data -> deixa em branco


def to_money(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return f"R$ {v:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    s = str(v).strip()
    if s in ("", "-", "?"):
        return ""
    return s


def to_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    s = str(v).strip()
    return "" if s == "-" else s


def main():
    xlsx_path = sys.argv[1]
    out_path = sys.argv[2]

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    categorias = set()
    contas = []

    for mes in MESES:
        ws = wb[mes]
        for row in ws.iter_rows(min_row=4, max_row=ws.max_row, max_col=9):
            categoria = to_text(row[0].value)
            if not categoria:
                continue
            categoria = normalize_categoria(categoria)
            beneficiario = to_text(row[1].value)
            identificacao = to_text(row[2].value)
            vencimento = to_iso_date(row[3].value)
            data_debito = to_iso_date(row[4].value)
            saida = to_money(row[5].value)
            status_raw = to_text(row[7].value).upper()
            status = STATUS_MAP.get(status_raw, "Em Aberto")
            observacoes = to_text(row[8].value)

            categorias.add(categoria)
            contas.append((categoria, beneficiario, identificacao, vencimento, data_debito, saida, status, observacoes))

    lines = []
    lines.append("-- =====================================================================")
    lines.append("-- Migration: importação dos lançamentos de CTR03.R00_FINANCEIRO_2026.xlsx (JAN..OUT)")
    lines.append("-- Gerado automaticamente por scripts/import-contas-xlsx.py")
    lines.append("-- Idempotente: só insere se a tabela `contas` estiver vazia.")
    lines.append("-- =====================================================================")
    lines.append("")
    lines.append("insert into public.categorias (nome)")
    lines.append("values")
    lines.append(",\n".join(f"  ('{esc(c)}')" for c in sorted(categorias)))
    lines.append("on conflict (nome) do nothing;")
    lines.append("")
    lines.append("do $$")
    lines.append("begin")
    lines.append("  if (select count(*) from public.contas) = 0 then")
    lines.append("    insert into public.contas (categoria, beneficiario, identificacao, vencimento, data_debito, saida, status, observacoes) values")
    row_sql = []
    for (categoria, beneficiario, identificacao, vencimento, data_debito, saida, status, observacoes) in contas:
        vals = [categoria, beneficiario, identificacao, vencimento, data_debito, saida, status, observacoes]
        sql_vals = ", ".join(f"'{esc(v)}'" for v in vals)
        row_sql.append(f"    ({sql_vals})")
    lines.append(",\n".join(row_sql) + ";")
    lines.append("  end if;")
    lines.append("end $$;")
    lines.append("")

    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    print(f"OK: {len(contas)} contas, {len(categorias)} categorias -> {out_path}")


if __name__ == "__main__":
    main()
